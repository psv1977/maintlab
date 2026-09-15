from datetime import date, datetime

from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin, UserPassesTestMixin
from django.db import transaction
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views import View
from django.views.generic import CreateView, DetailView, ListView, UpdateView

from .forms import EquipmentForm, EquipmentIdentifierForm, EquipmentImportForm, LocationForm, MeterReadingForm
from .models import Equipment, EquipmentIdentifier, Location, MeterReading
from organizations.tenant import get_user_organization


class EquipmentListView(LoginRequiredMixin, ListView):
    model = Equipment
    template_name = "equipment/equipment_list.html"
    context_object_name = "equipments"
    paginate_by = 20

    def get_queryset(self):
        queryset = super().get_queryset().filter(organization=get_user_organization(self.request.user))
        query = self.request.GET.get("q", "").strip()
        if query:
            queryset = queryset.filter(
                Q(code__icontains=query)
                | Q(name__icontains=query)
                | Q(serial_number__icontains=query)
            )

        status = self.request.GET.get("status")
        valid_statuses = {value for value, _ in Equipment.Status.choices}
        if status in valid_statuses:
            queryset = queryset.filter(status=status)

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        query_params = self.request.GET.copy()
        query_params.pop("page", None)
        context["query_params"] = query_params.urlencode()
        context["status_choices"] = Equipment.Status.choices
        return context


class EquipmentDetailView(LoginRequiredMixin, DetailView):
    model = Equipment
    template_name = "equipment/equipment_detail.html"
    context_object_name = "equipment"

    def get_queryset(self):
        return super().get_queryset().filter(organization=get_user_organization(self.request.user))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["meter_readings"] = self.object.meter_readings.select_related("recorded_by")[:10]
        context["latest_meter_reading"] = context["meter_readings"][0] if context["meter_readings"] else None
        context["maintenance_plans"] = self.object.maintenance_plans.filter(active=True)
        context["identifiers"] = self.object.identifiers.all()
        return context


class EquipmentCreateView(LoginRequiredMixin, CreateView):
    model = Equipment
    form_class = EquipmentForm
    template_name = "equipment/equipment_form.html"

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["organization"] = get_user_organization(self.request.user)
        return kwargs

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        form.instance.organization = get_user_organization(self.request.user)
        form.instance.status = Equipment.Status.OPERATIONAL
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("equipment:list")


class LocationCreateView(LoginRequiredMixin, CreateView):
    model = Location
    form_class = LocationForm
    template_name = "equipment/location_form.html"

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["organization"] = get_user_organization(self.request.user)
        return kwargs

    def form_valid(self, form):
        form.instance.organization = get_user_organization(self.request.user)
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("equipment:create")


class MeterReadingCreateView(LoginRequiredMixin, CreateView):
    model = MeterReading
    form_class = MeterReadingForm
    template_name = "equipment/meter_reading_form.html"

    def dispatch(self, request, *args, **kwargs):
        self.equipment = get_object_or_404(
            Equipment,
            pk=kwargs["equipment_id"],
            organization=get_user_organization(request.user),
        )
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["equipment"] = self.equipment
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["equipment"] = self.equipment
        return context

    def form_valid(self, form):
        form.instance.equipment = self.equipment
        form.instance.organization = self.equipment.organization
        form.instance.recorded_by = self.request.user
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("equipment:detail", args=[self.equipment.pk])


class EquipmentIdentifierCreateView(LoginRequiredMixin, CreateView):
    model = EquipmentIdentifier
    form_class = EquipmentIdentifierForm
    template_name = "equipment/equipment_identifier_form.html"

    def dispatch(self, request, *args, **kwargs):
        self.equipment = get_object_or_404(
            Equipment, pk=kwargs["equipment_id"], organization=get_user_organization(request.user),
        )
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["equipment"] = self.equipment
        return context

    def form_valid(self, form):
        form.instance.equipment = self.equipment
        form.instance.organization = self.equipment.organization
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("equipment:detail", args=[self.equipment.pk])


class EquipmentUpdateView(LoginRequiredMixin, UpdateView):
    model = Equipment
    form_class = EquipmentForm
    template_name = "equipment/equipment_form.html"

    def get_queryset(self):
        return super().get_queryset().filter(organization=get_user_organization(self.request.user))

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["organization"] = get_user_organization(self.request.user)
        return kwargs

    def form_valid(self, form):
        if form.instance.status == Equipment.Status.RETIRED:
            form.add_error("status", "No puede asignar el estado Retirado.")
            return self.form_invalid(form)
        form.instance.updated_by = self.request.user
        form.instance.updated_at = timezone.now()
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("equipment:detail", args=[self.object.pk])


class EquipmentRetireView(PermissionRequiredMixin, View):
    permission_required = "equipment.retire_equipment"
    template_name = "equipment/equipment_confirm_retire.html"

    def get(self, request, pk):
        equipment = get_object_or_404(Equipment, pk=pk, organization=get_user_organization(request.user))
        return render(request, self.template_name, {"equipment": equipment})

    def post(self, request, pk):
        equipment = get_object_or_404(Equipment, pk=pk, organization=get_user_organization(request.user))
        equipment.status = Equipment.Status.RETIRED
        equipment.save(update_fields=["status"])
        return redirect(reverse("equipment:detail", args=[equipment.pk]))


class EquipmentImportView(LoginRequiredMixin, UserPassesTestMixin, View):
    template_name = "equipment/equipment_import.html"

    def test_func(self):
        return self.request.user.is_staff

    def get(self, request):
        return render(request, self.template_name, {"form": EquipmentImportForm()})

    def post(self, request):
        form = EquipmentImportForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {"form": form})

        uploaded_file = form.cleaned_data["file"]
        try:
            if uploaded_file.name.lower().endswith(".csv"):
                headers, data_rows = self._parse_csv(uploaded_file)
            else:
                headers, data_rows = self._parse_xlsx(uploaded_file)
        except (UnicodeDecodeError, ValueError):
            form.add_error("file", "No fue posible leer el archivo seleccionado.")
            return render(request, self.template_name, {"form": form})

        required_headers = ["name", "code"]
        if any(header not in headers for header in required_headers):
            form.add_error("file", "La planilla debe incluir las columnas name y code.")
            return render(request, self.template_name, {"form": form})

        organization = get_user_organization(request.user)
        rows = []
        errors = []
        seen_codes = set()
        for row_number, values in enumerate(data_rows, start=2):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            data = dict(zip(headers, values))
            name = str(data.get("name") or "").strip()
            code = str(data.get("code") or "").strip()
            if not name or not code:
                errors.append(f"Fila {row_number}: name y code son obligatorios.")
            elif code in seen_codes or Equipment.objects.filter(organization=organization, code=code).exists():
                errors.append(f"Fila {row_number}: el código {code} ya existe.")
            else:
                try:
                    data["commissioned_at"] = self._parse_date(data.get("commissioned_at"))
                except ValueError:
                    errors.append(f"Fila {row_number}: commissioned_at debe usar el formato YYYY-MM-DD.")
                    continue
                seen_codes.add(code)
                rows.append(data)

        if errors:
            return render(request, self.template_name, {"form": form, "errors": errors})

        valid_types = {value for value, _ in Equipment.EquipmentType.choices}
        valid_statuses = {value for value, _ in Equipment.Status.choices}

        with transaction.atomic():
            for data in rows:
                location_name = str(data.get("location") or "").strip()
                location = None
                if location_name:
                    location, _ = Location.objects.get_or_create(organization=organization, name=location_name)

                equipment_type = str(data.get("equipment_type") or "").strip().lower()
                if equipment_type not in valid_types:
                    equipment_type = Equipment.EquipmentType.INDUSTRIAL

                status = str(data.get("status") or "").strip().lower()
                if status not in valid_statuses:
                    status = Equipment.Status.OPERATIONAL

                Equipment.objects.create(
                    organization=organization,
                    name=str(data["name"]).strip(),
                    code=str(data["code"]).strip(),
                    equipment_type=equipment_type,
                    description=str(data.get("description") or "").strip(),
                    serial_number=str(data.get("serial_number") or "").strip(),
                    brand=str(data.get("brand") or "").strip(),
                    model=str(data.get("model") or "").strip(),
                    location=location,
                    commissioned_at=data["commissioned_at"],
                    application=str(data.get("application") or "").strip(),
                    status=status,
                    created_by=request.user,
                )
        return render(request, self.template_name, {"form": EquipmentImportForm(), "imported_count": len(rows)})

    def _parse_csv(self, uploaded_file):
        import csv
        import io

        text = uploaded_file.read().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return [], []
        headers = [h.strip().lower() for h in rows[0]]
        return headers, rows[1:]

    def _parse_xlsx(self, uploaded_file):
        from openpyxl import load_workbook

        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        all_rows = list(worksheet.iter_rows(values_only=True))
        if not all_rows:
            return [], []
        headers = [str(v).strip().lower() if v is not None else "" for v in all_rows[0]]
        return headers, all_rows[1:]

    def _parse_date(self, value):
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
